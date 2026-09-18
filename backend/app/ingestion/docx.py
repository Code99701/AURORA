import re
from io import BytesIO
from typing import BinaryIO

from docx import Document as OpenDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.ingestion.base import (
    AdapterCapabilities,
    AdapterError,
    ContentUnitDraft,
    ExtractionContext,
)
from app.ingestion.common import finalize_document, normalized_text, read_bounded
from app.ingestion.detection import inspect_ooxml_archive


HEADING_STYLE = re.compile(r"heading\s+(\d+)", re.IGNORECASE)


def _paragraph_metadata(paragraph: Paragraph) -> dict:
    style_name = paragraph.style.name if paragraph.style is not None else None
    numbering = False
    try:
        numbering = paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
    except AttributeError:
        pass
    return {
        "style": style_name,
        "is_list_item": numbering
        or bool(style_name and style_name.casefold().startswith("list")),
    }


def _render_table(
    table: Table, remaining_cells: int, *, depth: int = 0
) -> tuple[str, int, int]:
    rows: list[str] = []
    cells_seen = 0
    nested_tables = 0
    for row_index, row in enumerate(table.rows, start=1):
        values: list[str] = []
        for column_index, cell in enumerate(row.cells, start=1):
            cells_seen += 1
            if cells_seen > remaining_cells:
                raise AdapterError(
                    "TABULAR_CELL_LIMIT",
                    "Document tables exceed the configured cell limit",
                    422,
                    failure_stage="EXTRACTION",
                )
            value = normalized_text(cell.text)
            values.append(value)
            for nested_index, nested_table in enumerate(cell.tables, start=1):
                nested_text, nested_cells, nested_count = _render_table(
                    nested_table, remaining_cells - cells_seen, depth=depth + 1
                )
                cells_seen += nested_cells
                nested_tables += nested_count + 1
                if nested_text:
                    rows.append(
                        f"Nested table at row {row_index}, column {column_index}, "
                        f"depth {depth + 1}, table {nested_index}:\n{nested_text}"
                    )
        if any(values):
            rows.append(
                f"Row {row_index}: "
                + " | ".join(
                    f"Column {column_index}: {value}"
                    for column_index, value in enumerate(values, start=1)
                    if value
                )
            )
    return "\n".join(rows), cells_seen, nested_tables


class DOCXAdapter:
    adapter_id = "aurora.docx"
    adapter_version = "1.0"
    source_format = "docx"
    capabilities = AdapterCapabilities(
        extensions=(".docx",),
        media_types=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        unit_kinds=("HEADING", "PARAGRAPH", "TABLE", "HEADER", "FOOTER"),
        supports_partial_results=True,
        supports_native_extraction=True,
        supports_ocr_fallback=False,
    )

    def extract(self, source: BinaryIO, context: ExtractionContext):
        data = read_bounded(source, context.limits.maximum_file_bytes)
        try:
            archive = inspect_ooxml_archive(
                BytesIO(data),
                required_member="word/document.xml",
                maximum_entries=context.limits.maximum_archive_entries,
                maximum_uncompressed_bytes=context.limits.maximum_archive_uncompressed_bytes,
            )
            document = OpenDocument(BytesIO(data))
        except (PackageNotFoundError, ValueError, KeyError, OSError) as exc:
            raise AdapterError(
                "DOCX_PARSE_ERROR",
                "Word document could not be parsed safely",
                422,
                failure_stage="PARSING",
            ) from exc
        except Exception as exc:
            raise AdapterError(
                "DOCX_PARSE_ERROR",
                "Word document could not be parsed safely",
                422,
                failure_stage="PARSING",
            ) from exc

        units: list[ContentUnitDraft] = []
        heading_ordinals: dict[int, int] = {}
        heading_path: list[str] = []
        paragraph_index = 0
        table_index = 0
        paragraph_count = 0
        table_cell_count = 0
        nested_table_count = 0
        failed_items = 0
        empty_items = 0

        try:
            blocks = document.iter_inner_content()
            for block in blocks:
                if isinstance(block, Paragraph):
                    paragraph_index += 1
                    paragraph_count += 1
                    if paragraph_count > context.limits.maximum_paragraphs:
                        raise AdapterError(
                            "DOCX_PARAGRAPH_LIMIT",
                            "Word document exceeds the configured paragraph limit",
                            422,
                            failure_stage="EXTRACTION",
                        )
                    try:
                        text = normalized_text(block.text)
                        metadata = _paragraph_metadata(block)
                    except Exception:
                        failed_items += 1
                        units.append(
                            ContentUnitDraft(
                                ordinal=len(units),
                                kind="PARAGRAPH",
                                text="",
                                locator={
                                    "v": 1,
                                    "type": "docx_paragraph",
                                    "paragraph": paragraph_index,
                                },
                                extraction_method="NATIVE_FAILED",
                                content_type=self.capabilities.media_types[0],
                                status="FAILED",
                                issue_code="DOCX_PARAGRAPH_ERROR",
                            )
                        )
                        continue
                    if not text:
                        empty_items += 1
                        continue
                    style_name = metadata["style"] or ""
                    heading_match = HEADING_STYLE.fullmatch(style_name.strip())
                    if heading_match:
                        level = max(1, min(9, int(heading_match.group(1))))
                        heading_path = heading_path[: level - 1]
                        heading_path.append(text)
                        parent_ordinal = next(
                            (
                                heading_ordinals[parent]
                                for parent in range(level - 1, 0, -1)
                                if parent in heading_ordinals
                            ),
                            None,
                        )
                        heading_ordinals = {
                            key: value for key, value in heading_ordinals.items() if key < level
                        }
                        ordinal = len(units)
                        units.append(
                            ContentUnitDraft(
                                ordinal=ordinal,
                                kind="HEADING",
                                text=text,
                                locator={
                                    "v": 1,
                                    "type": "docx_paragraph",
                                    "paragraph": paragraph_index,
                                    "heading_path": list(heading_path),
                                },
                                extraction_method="NATIVE",
                                content_type=self.capabilities.media_types[0],
                                parent_ordinal=parent_ordinal,
                                metadata={**metadata, "heading_level": level},
                            )
                        )
                        heading_ordinals[level] = ordinal
                    else:
                        parent_ordinal = next(
                            (heading_ordinals[level] for level in sorted(heading_ordinals, reverse=True)),
                            None,
                        )
                        units.append(
                            ContentUnitDraft(
                                ordinal=len(units),
                                kind="PARAGRAPH",
                                text=text,
                                locator={
                                    "v": 1,
                                    "type": "docx_paragraph",
                                    "paragraph": paragraph_index,
                                    "heading_path": list(heading_path),
                                },
                                extraction_method="NATIVE",
                                content_type=self.capabilities.media_types[0],
                                parent_ordinal=parent_ordinal,
                                metadata=metadata,
                            )
                        )
                elif isinstance(block, Table):
                    table_index += 1
                    try:
                        text, cells, nested = _render_table(
                            block, context.limits.maximum_table_cells - table_cell_count
                        )
                        table_cell_count += cells
                        nested_table_count += nested
                    except AdapterError:
                        raise
                    except Exception:
                        failed_items += 1
                        units.append(
                            ContentUnitDraft(
                                ordinal=len(units),
                                kind="TABLE",
                                text="",
                                locator={"v": 1, "type": "docx_table", "table": table_index},
                                extraction_method="NATIVE_FAILED",
                                content_type=self.capabilities.media_types[0],
                                status="FAILED",
                                issue_code="DOCX_TABLE_ERROR",
                            )
                        )
                        continue
                    if not text:
                        empty_items += 1
                        continue
                    parent_ordinal = next(
                        (heading_ordinals[level] for level in sorted(heading_ordinals, reverse=True)),
                        None,
                    )
                    units.append(
                        ContentUnitDraft(
                            ordinal=len(units),
                            kind="TABLE",
                            text=text,
                            locator={
                                "v": 1,
                                "type": "docx_table",
                                "table": table_index,
                                "heading_path": list(heading_path),
                            },
                            extraction_method="NATIVE",
                            content_type=self.capabilities.media_types[0],
                            parent_ordinal=parent_ordinal,
                            metadata={
                                "cell_count": cells,
                                "row_count": len(block.rows),
                                "column_count": len(block.columns),
                                "nested_table_count": nested,
                            },
                        )
                    )

            seen_parts: set[tuple[str, str]] = set()
            for section_index, section in enumerate(document.sections, start=1):
                for kind, container in (("HEADER", section.header), ("FOOTER", section.footer)):
                    key = (kind, str(container.part.partname))
                    if key in seen_parts:
                        continue
                    seen_parts.add(key)
                    text = normalized_text("\n".join(p.text for p in container.paragraphs))
                    if not text:
                        continue
                    units.append(
                        ContentUnitDraft(
                            ordinal=len(units),
                            kind=kind,
                            text=text,
                            locator={
                                "v": 1,
                                "type": f"docx_{kind.casefold()}",
                                "section": section_index,
                            },
                            extraction_method="NATIVE",
                            content_type=self.capabilities.media_types[0],
                            metadata={"linked_to_previous": container.is_linked_to_previous},
                        )
                    )
        except AdapterError:
            raise
        except Exception as exc:
            raise AdapterError(
                "DOCX_PARSE_ERROR",
                "Word document could not be parsed safely",
                422,
                failure_stage="EXTRACTION",
            ) from exc

        return finalize_document(
            source_format=self.source_format,
            adapter_id=self.adapter_id,
            adapter_version=self.adapter_version,
            units=units,
            limits=context.limits,
            diagnostics={
                **archive,
                "paragraph_count": paragraph_count,
                "table_count": table_index,
                "table_cell_count": table_cell_count,
                "nested_table_count": nested_table_count,
                "empty_item_count": empty_items,
                "failed_item_count": failed_items,
                "inline_image_count": len(document.inline_shapes),
                "section_count": len(document.sections),
            },
        )
