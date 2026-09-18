from io import BytesIO
from typing import BinaryIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.ingestion.base import (
    AdapterCapabilities,
    AdapterError,
    ContentUnitDraft,
    ExtractionContext,
)
from app.ingestion.common import finalize_document, normalized_text, read_bounded
from app.ingestion.detection import inspect_ooxml_archive


class PPTXAdapter:
    adapter_id = "aurora.pptx"
    adapter_version = "1.0"
    source_format = "pptx"
    capabilities = AdapterCapabilities(
        extensions=(".pptx",),
        media_types=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
        unit_kinds=("SLIDE",),
        supports_partial_results=True,
        supports_native_extraction=True,
        supports_ocr_fallback=False,
    )

    def extract(self, source: BinaryIO, context: ExtractionContext):
        data = read_bounded(source, context.limits.maximum_file_bytes)
        try:
            archive = inspect_ooxml_archive(
                BytesIO(data),
                required_member="ppt/presentation.xml",
                maximum_entries=context.limits.maximum_archive_entries,
                maximum_uncompressed_bytes=context.limits.maximum_archive_uncompressed_bytes,
            )
            presentation = Presentation(BytesIO(data))
        except (ValueError, KeyError, OSError) as exc:
            raise AdapterError(
                "PPTX_PARSE_ERROR",
                "PowerPoint presentation could not be parsed safely",
                422,
                failure_stage="PARSING",
            ) from exc
        except Exception as exc:
            raise AdapterError(
                "PPTX_PARSE_ERROR",
                "PowerPoint presentation could not be parsed safely",
                422,
                failure_stage="PARSING",
            ) from exc

        if len(presentation.slides) > context.limits.maximum_slides:
            raise AdapterError(
                "PPTX_SLIDE_LIMIT",
                "Presentation exceeds the configured slide limit",
                422,
                failure_stage="EXTRACTION",
            )

        units: list[ContentUnitDraft] = []
        total_table_cells = 0
        image_count = 0
        chart_count = 0
        table_count = 0
        notes_slide_count = 0
        failed_slides = 0

        for slide_number, slide in enumerate(presentation.slides, start=1):
            try:
                lines: list[str] = []
                shape_metadata: list[dict] = []
                slide_images = 0
                slide_charts = 0
                slide_tables = 0
                for shape_order, shape in enumerate(slide.shapes, start=1):
                    descriptor = {
                        "order": shape_order,
                        "shape_type": str(shape.shape_type),
                        "has_text": bool(getattr(shape, "has_text_frame", False)),
                        "has_table": bool(getattr(shape, "has_table", False)),
                        "has_chart": bool(getattr(shape, "has_chart", False)),
                    }
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        slide_images += 1
                    if getattr(shape, "has_chart", False):
                        slide_charts += 1
                    if getattr(shape, "has_text_frame", False):
                        for paragraph_order, paragraph in enumerate(
                            shape.text_frame.paragraphs, start=1
                        ):
                            text = normalized_text(paragraph.text)
                            if text:
                                prefix = "  " * max(0, int(paragraph.level or 0))
                                lines.append(
                                    f"Shape {shape_order}, paragraph {paragraph_order}: {prefix}{text}"
                                )
                    if getattr(shape, "has_table", False):
                        slide_tables += 1
                        rendered_rows: list[str] = []
                        for row_number, row in enumerate(shape.table.rows, start=1):
                            values = []
                            for column_number, cell in enumerate(row.cells, start=1):
                                total_table_cells += 1
                                if total_table_cells > context.limits.maximum_table_cells:
                                    raise AdapterError(
                                        "TABULAR_CELL_LIMIT",
                                        "Presentation tables exceed the configured cell limit",
                                        422,
                                        failure_stage="EXTRACTION",
                                    )
                                value = normalized_text(cell.text)
                                if value:
                                    values.append(f"Column {column_number}: {value}")
                            if values:
                                rendered_rows.append(
                                    f"Table shape {shape_order}, row {row_number}: "
                                    + " | ".join(values)
                                )
                        lines.extend(rendered_rows)
                    shape_metadata.append(descriptor)

                notes_text = normalized_text(slide.notes_slide.notes_text_frame.text)
                if notes_text:
                    notes_slide_count += 1
                    lines.append(f"Speaker notes: {notes_text}")
                text = "\n".join(lines)
                image_count += slide_images
                chart_count += slide_charts
                table_count += slide_tables
                title = normalized_text(slide.shapes.title.text) if slide.shapes.title else None
                metadata = {
                    "title": title,
                    "shape_count": len(slide.shapes),
                    "image_count": slide_images,
                    "chart_count": slide_charts,
                    "table_count": slide_tables,
                    "has_speaker_notes": bool(notes_text),
                    "shapes": shape_metadata,
                }
                if text:
                    units.append(
                        ContentUnitDraft(
                            ordinal=len(units),
                            kind="SLIDE",
                            text=text,
                            locator={
                                "v": 1,
                                "type": "pptx_slide",
                                "slide": slide_number,
                            },
                            extraction_method="NATIVE",
                            content_type=self.capabilities.media_types[0],
                            metadata=metadata,
                        )
                    )
                else:
                    issue_code = (
                        "VISUAL_ONLY_SLIDE"
                        if slide_images or slide_charts
                        else "EMPTY_SLIDE"
                    )
                    units.append(
                        ContentUnitDraft(
                            ordinal=len(units),
                            kind="SLIDE",
                            text="",
                            locator={
                                "v": 1,
                                "type": "pptx_slide",
                                "slide": slide_number,
                            },
                            extraction_method="NATIVE",
                            content_type=self.capabilities.media_types[0],
                            status="SKIPPED",
                            issue_code=issue_code,
                            metadata=metadata,
                        )
                    )
            except AdapterError:
                raise
            except Exception:
                failed_slides += 1
                units.append(
                    ContentUnitDraft(
                        ordinal=len(units),
                        kind="SLIDE",
                        text="",
                        locator={"v": 1, "type": "pptx_slide", "slide": slide_number},
                        extraction_method="NATIVE_FAILED",
                        content_type=self.capabilities.media_types[0],
                        status="FAILED",
                        issue_code="PPTX_SLIDE_ERROR",
                    )
                )

        return finalize_document(
            source_format=self.source_format,
            adapter_id=self.adapter_id,
            adapter_version=self.adapter_version,
            units=units,
            limits=context.limits,
            diagnostics={
                **archive,
                "slide_count": len(presentation.slides),
                "failed_slide_count": failed_slides,
                "image_count": image_count,
                "chart_count": chart_count,
                "table_count": table_count,
                "table_cell_count": total_table_cells,
                "speaker_notes_slide_count": notes_slide_count,
                "visual_content_interpreted": False,
            },
        )

