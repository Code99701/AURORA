from io import BytesIO
from pathlib import Path
from uuid import UUID

import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from pptx import Presentation
from pptx.util import Inches
from zipfile import ZIP_DEFLATED, ZipFile

from app.ingestion.base import AdapterError, ExtractionContext, ExtractionLimits
from app.ingestion.docx import DOCXAdapter
from app.ingestion.pptx import PPTXAdapter
from app.ingestion.registry import get_adapter_registry
from app.ingestion.tabular import CSVAdapter, TSVAdapter
from app.ingestion.text import MarkdownAdapter, TXTAdapter
from app.ingestion.xlsx import XLSXAdapter
from app.embeddings.deps import get_embedding_provider
from app.llm import GenerationResult
from app.llm.deps import get_llm_provider
from app.main import app
from app.models.content_unit import ContentUnit
from app.models.document import Document
from app.models.document_chunk import EMBEDDING_DIMENSION
from app.models.ingestion_run import IngestionRun
from app.ocr.deps import get_ocr_provider
from app.storage.deps import get_file_storage
from app.storage.local import LocalFileStorage
from app.db.session import SessionLocal
from sqlalchemy import select
from tests.adapter_compliance import (
    assert_adapter_capabilities,
    assert_canonical_document_contract,
)
from tests.test_processing import build_text_pdf


@pytest.fixture
def multiformat_limits() -> ExtractionLimits:
    return ExtractionLimits(
        maximum_file_bytes=2_000_000,
        maximum_pages=10,
        render_dpi=72,
        maximum_rendered_pixels_per_page=1_000_000,
        maximum_total_rendered_pixels=2_000_000,
        maximum_ocr_pages=5,
        ocr_timeout_seconds_per_page=3,
        native_minimum_characters=8,
        ocr_minimum_characters=12,
        minimum_printable_ratio=0.85,
        minimum_ocr_confidence=0.25,
        maximum_extracted_characters=100_000,
        maximum_units=1_000,
        maximum_rows=1_000,
        maximum_columns=100,
        maximum_table_cells=10_000,
        maximum_field_characters=10_000,
        maximum_archive_entries=2_000,
        maximum_archive_uncompressed_bytes=10_000_000,
        maximum_paragraphs=1_000,
        maximum_slides=100,
        maximum_sheets=50,
        tabular_rows_per_unit=2,
    )


def context(limits: ExtractionLimits) -> ExtractionContext:
    return ExtractionContext(ocr_provider=object(), limits=limits)


def word_bytes() -> bytes:
    document = WordDocument()
    document.add_heading("Aurora evidence", level=1)
    document.add_paragraph("नमस्ते — multilingual paragraph")
    document.add_paragraph("First research item", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "Verified"
    document.sections[0].header.paragraphs[0].text = "Private research header"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def workbook_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Research Data"
    sheet.append(["Topic", "Value", "Formula"])
    sheet.append(["database", 42, "=B2*2"])
    sheet.append(["Unicode", "हिमालय", True])
    report_table = Table(displayName="ResearchTable", ref="A1:C3")
    report_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2")
    sheet.add_table(report_table)
    sheet.merge_cells("A4:B4")
    sheet["A4"] = "Merged conclusion"
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions["C"].hidden = True
    hidden = workbook.create_sheet("Hidden evidence")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "restricted note"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def presentation_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Research findings"
    slide.placeholders[1].text = "Database evidence\nMultilingual नमस्ते"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Coverage"
    table.cell(1, 1).text = "Complete"
    slide.notes_slide.notes_text_frame.text = "Speaker-only supporting evidence"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def package_with_active_content(data: bytes, member: str) -> bytes:
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry, source.read(entry.filename))
        target.writestr(member, b"hostile active payload")
    return output.getvalue()


def empty_word_bytes() -> bytes:
    output = BytesIO()
    WordDocument().save(output)
    return output.getvalue()


def empty_workbook_bytes() -> bytes:
    output = BytesIO()
    Workbook().save(output)
    return output.getvalue()


def empty_presentation_bytes() -> bytes:
    output = BytesIO()
    Presentation().save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "adapter",
    [
        TXTAdapter(),
        MarkdownAdapter(),
        CSVAdapter(),
        TSVAdapter(),
        DOCXAdapter(),
        XLSXAdapter(),
        PPTXAdapter(),
    ],
)
def test_all_new_adapters_declare_the_shared_contract(adapter):
    assert_adapter_capabilities(adapter)


def test_registry_resolves_every_supported_format_by_mime_and_extension():
    registry = get_adapter_registry()
    expected = {
        ("text/plain", "notes.txt"): "txt",
        ("text/markdown", "notes.md"): "markdown",
        ("text/csv", "data.csv"): "csv",
        ("text/tab-separated-values", "data.tsv"): "tsv",
        (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "report.docx",
        ): "docx",
        (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "report.xlsx",
        ): "xlsx",
        (
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "report.pptx",
        ): "pptx",
        ("application/pdf", "report.pdf"): "pdf",
    }
    assert {
        key: registry.resolve_document(*key).source_format for key in expected
    } == expected


def test_txt_handles_bom_unicode_mixed_newlines_and_stable_ranges(multiformat_limits):
    source = b"\xef\xbb\xbfFirst line\r\n\r\nUnicode: \xe0\xa4\xa8\xe0\xa4\xae\xe0\xa4\xb8\xe0\xa5\x8d\xe0\xa4\xa4\xe0\xa5\x87\rLast line\n"
    result = TXTAdapter().extract(BytesIO(source), context(multiformat_limits))
    assert [unit.locator["line_start"] for unit in result.units] == [1, 3]
    assert "नमस्ते" in result.units[1].text
    assert_canonical_document_contract(result)


def test_txt_replaces_invalid_utf8_without_leaking_it_to_diagnostics(multiformat_limits):
    result = TXTAdapter().extract(
        BytesIO(b"trustworthy \xff evidence"), context(multiformat_limits)
    )
    assert result.warnings == ("INVALID_UTF8_REPLACED",)
    assert result.diagnostics["invalid_unicode_replaced"] is True
    assert "trustworthy" not in repr(result.diagnostics)


def test_markdown_preserves_nested_heading_provenance_and_inert_content(multiformat_limits):
    source = """# Study
Intro with [reference](https://example.invalid).

## Findings
- database result
> quoted evidence
```python
raise RuntimeError('must never execute')
```

| Metric | Value |
| --- | --- |
| Coverage | पूर्ण |
""".encode()
    result = MarkdownAdapter().extract(BytesIO(source), context(multiformat_limits))
    assert [unit.locator["heading_path"] for unit in result.units] == [
        ["Study"],
        ["Study", "Findings"],
    ]
    assert result.units[1].parent_ordinal == 0
    assert result.units[1].metadata["contains_code_fence"] is True
    assert result.units[0].metadata["link_count"] == 1
    assert_canonical_document_contract(result)


@pytest.mark.parametrize(
    ("adapter", "source", "delimiter"),
    [
        (
            CSVAdapter(),
            '\ufeffname,detail,value\r\nAurora,"line one\nline two",=1+1\r\nSparse,,नमस्ते\r\n',
            ",",
        ),
        (TSVAdapter(), "name\tdetail\nAurora\tquoted evidence\n", "\t"),
    ],
)
def test_delimited_formats_preserve_rows_headers_and_formula_text(
    multiformat_limits, adapter, source, delimiter
):
    result = adapter.extract(BytesIO(source.encode()), context(multiformat_limits))
    assert result.units[0].locator["type"] == "delimited_rows"
    assert result.units[0].metadata["delimiter"] == delimiter
    assert result.units[0].metadata["has_header"] is True
    if adapter.source_format == "csv":
        assert "=1+1" in result.units[0].text
        assert result.diagnostics["formula_like_value_count"] == 1
    assert_canonical_document_contract(result)


def test_docx_preserves_headings_paragraphs_lists_tables_and_headers(multiformat_limits):
    result = DOCXAdapter().extract(BytesIO(word_bytes()), context(multiformat_limits))
    assert [unit.kind for unit in result.units] == [
        "HEADING",
        "PARAGRAPH",
        "PARAGRAPH",
        "TABLE",
        "HEADER",
    ]
    assert result.units[1].parent_ordinal == 0
    assert result.units[2].metadata["is_list_item"] is True
    assert "Row 2: Column 1: Accuracy | Column 2: Verified" in result.units[3].text
    assert result.page_count is None
    assert_canonical_document_contract(result)


def test_xlsx_preserves_sheets_coordinates_formulas_types_and_hidden_state(multiformat_limits):
    result = XLSXAdapter().extract(BytesIO(workbook_bytes()), context(multiformat_limits))
    assert result.diagnostics["sheet_count"] == 2
    assert result.diagnostics["formula_count"] == 1
    assert result.diagnostics["formula_calculation_performed"] is False
    assert any(unit.locator["sheet"] == "Research Data" for unit in result.units)
    assert any(unit.metadata["sheet_state"] == "hidden" for unit in result.units)
    assert any("C2 [formula]: =B2*2" in unit.text for unit in result.units)
    research_units = [
        unit for unit in result.units if unit.locator["sheet"] == "Research Data"
    ]
    assert research_units[0].metadata["merged_ranges"] == ["A4:B4"]
    assert research_units[0].metadata["hidden_rows"] == [3]
    assert research_units[0].metadata["hidden_columns"] == [
        {"minimum": 3, "maximum": 3}
    ]
    assert result.diagnostics["defined_tables"] == [
        {"name": "ResearchTable", "display_name": "ResearchTable", "range": "A1:C3"}
    ]
    assert result.page_count is None
    assert_canonical_document_contract(result)


def test_pptx_preserves_slide_text_table_notes_and_empty_slide(multiformat_limits):
    result = PPTXAdapter().extract(
        BytesIO(presentation_bytes()), context(multiformat_limits)
    )
    assert result.units[0].locator == {"v": 1, "type": "pptx_slide", "slide": 1}
    assert "Speaker-only supporting evidence" in result.units[0].text
    assert result.units[0].metadata["table_count"] == 1
    assert result.units[1].status == "SKIPPED"
    assert result.units[1].issue_code == "EMPTY_SLIDE"
    assert result.page_count is None
    assert_canonical_document_contract(result)


@pytest.mark.parametrize(
    ("adapter", "source"),
    [
        (TXTAdapter(), b"   \r\n\t"),
        (MarkdownAdapter(), b"<!-- formatting only -->"),
        (CSVAdapter(), b",,\n,,\n"),
        (TSVAdapter(), b"\t\t\n"),
    ],
)
def test_empty_textual_sources_fail_with_a_sanitized_error(
    multiformat_limits, adapter, source
):
    with pytest.raises(AdapterError) as error:
        adapter.extract(BytesIO(source), context(multiformat_limits))
    assert error.value.code == "NO_EXTRACTABLE_TEXT"
    original = source.decode(errors="ignore").strip()
    if original:
        assert original not in repr(error.value.diagnostics)


@pytest.mark.parametrize(
    ("adapter", "source", "code"),
    [
        (DOCXAdapter(), b"PK\x03\x04not-a-word-package", "DOCX_PARSE_ERROR"),
        (XLSXAdapter(), b"PK\x03\x04not-a-workbook", "XLSX_PARSE_ERROR"),
        (PPTXAdapter(), b"PK\x03\x04not-a-presentation", "PPTX_PARSE_ERROR"),
    ],
)
def test_malformed_ooxml_is_rejected_safely(multiformat_limits, adapter, source, code):
    with pytest.raises(AdapterError) as error:
        adapter.extract(BytesIO(source), context(multiformat_limits))
    assert error.value.code == code
    assert "not-a" not in error.value.public_message


@pytest.mark.parametrize(
    ("adapter", "source_factory"),
    [
        (DOCXAdapter(), empty_word_bytes),
        (XLSXAdapter(), empty_workbook_bytes),
        (PPTXAdapter(), empty_presentation_bytes),
    ],
)
def test_empty_ooxml_sources_do_not_claim_success(
    multiformat_limits, adapter, source_factory
):
    with pytest.raises(AdapterError) as error:
        adapter.extract(BytesIO(source_factory()), context(multiformat_limits))
    assert error.value.code == "NO_EXTRACTABLE_TEXT"
    assert error.value.diagnostics["usable_unit_count"] == 0


@pytest.mark.parametrize(
    ("adapter", "source_factory"),
    [
        (TXTAdapter(), lambda: b"First block\n\nSecond block"),
        (MarkdownAdapter(), lambda: b"# First\nEvidence\n## Second\nMore evidence"),
        (CSVAdapter(), lambda: b"name,value\nfirst,1\nsecond,2\n"),
        (TSVAdapter(), lambda: b"name\tvalue\nfirst\t1\nsecond\t2\n"),
        (DOCXAdapter(), word_bytes),
        (XLSXAdapter(), workbook_bytes),
        (PPTXAdapter(), presentation_bytes),
    ],
)
def test_new_adapters_are_deterministic(multiformat_limits, adapter, source_factory):
    source = source_factory()
    first = adapter.extract(BytesIO(source), context(multiformat_limits))
    second = adapter.extract(BytesIO(source), context(multiformat_limits))
    assert first == second


def test_office_structure_limits_fail_globally(multiformat_limits):
    word_limit = ExtractionLimits(
        **{**multiformat_limits.__dict__, "maximum_paragraphs": 1}
    )
    with pytest.raises(AdapterError) as word_error:
        DOCXAdapter().extract(BytesIO(word_bytes()), context(word_limit))
    assert word_error.value.code == "DOCX_PARAGRAPH_LIMIT"

    sheet_limit = ExtractionLimits(
        **{**multiformat_limits.__dict__, "maximum_sheets": 1}
    )
    with pytest.raises(AdapterError) as sheet_error:
        XLSXAdapter().extract(BytesIO(workbook_bytes()), context(sheet_limit))
    assert sheet_error.value.code == "XLSX_SHEET_LIMIT"

    slide_limit = ExtractionLimits(
        **{**multiformat_limits.__dict__, "maximum_slides": 1}
    )
    with pytest.raises(AdapterError) as slide_error:
        PPTXAdapter().extract(BytesIO(presentation_bytes()), context(slide_limit))
    assert slide_error.value.code == "PPTX_SLIDE_LIMIT"


def test_resource_limits_are_enforced_at_shared_boundaries(multiformat_limits):
    constrained = ExtractionLimits(
        **{**multiformat_limits.__dict__, "maximum_rows": 1}
    )
    with pytest.raises(AdapterError) as error:
        CSVAdapter().extract(BytesIO(b"name,value\na,1\nb,2\n"), context(constrained))
    assert error.value.code == "TABULAR_ROW_LIMIT"


def test_malformed_csv_is_rejected_without_parser_details(multiformat_limits):
    with pytest.raises(AdapterError) as error:
        CSVAdapter().extract(
            BytesIO(b'name,detail\nAurora,"unterminated\n'),
            context(multiformat_limits),
        )
    assert error.value.code == "DELIMITED_TEXT_PARSE_ERROR"
    assert "unterminated" not in error.value.public_message


def test_ooxml_active_content_is_rejected_before_parser_execution(multiformat_limits):
    hostile = package_with_active_content(word_bytes(), "word/vbaProject.bin")
    with pytest.raises(AdapterError) as error:
        DOCXAdapter().extract(BytesIO(hostile), context(multiformat_limits))
    assert error.value.code == "DOCX_PARSE_ERROR"


def test_docx_recovers_from_one_paragraph_failure(
    multiformat_limits, monkeypatch
):
    import app.ingestion.docx as docx_module

    document = WordDocument()
    document.add_paragraph("Usable paragraph")
    document.add_paragraph("Broken paragraph")
    output = BytesIO()
    document.save(output)
    original = docx_module.normalized_text

    def fail_one(value):
        if value == "Broken paragraph":
            raise ValueError("private parser detail")
        return original(value)

    monkeypatch.setattr(docx_module, "normalized_text", fail_one)
    result = DOCXAdapter().extract(BytesIO(output.getvalue()), context(multiformat_limits))
    assert result.outcome == "PARTIAL"
    assert [unit.status for unit in result.units] == ["USABLE", "FAILED"]
    assert result.units[1].issue_code == "DOCX_PARAGRAPH_ERROR"
    assert "private parser detail" not in repr(result.diagnostics)
    assert_canonical_document_contract(result)


def test_xlsx_recovers_from_one_worksheet_failure(
    multiformat_limits, monkeypatch
):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    workbook = Workbook()
    workbook.active.title = "Usable"
    workbook.active["A1"] = "Trustworthy evidence"
    workbook.create_sheet("Broken")["A1"] = "Broken private cell"
    output = BytesIO()
    workbook.save(output)
    original = ReadOnlyWorksheet.iter_rows

    def fail_one(self, *args, **kwargs):
        if self.title == "Broken":
            raise ValueError("private worksheet detail")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", fail_one)
    result = XLSXAdapter().extract(BytesIO(output.getvalue()), context(multiformat_limits))
    assert result.outcome == "PARTIAL"
    assert [unit.status for unit in result.units] == ["USABLE", "FAILED"]
    assert result.units[1].issue_code == "XLSX_WORKSHEET_ERROR"
    assert "private worksheet detail" not in repr(result.diagnostics)
    assert_canonical_document_contract(result)


def test_pptx_recovers_from_one_slide_failure(
    multiformat_limits, monkeypatch
):
    import app.ingestion.pptx as pptx_module

    presentation = Presentation()
    broken = presentation.slides.add_slide(presentation.slide_layouts[1])
    broken.shapes.title.text = "Broken slide"
    usable = presentation.slides.add_slide(presentation.slide_layouts[1])
    usable.shapes.title.text = "Usable slide"
    usable.placeholders[1].text = "Trustworthy evidence"
    output = BytesIO()
    presentation.save(output)
    original = pptx_module.normalized_text

    def fail_one(value):
        if value == "Broken slide":
            raise ValueError("private slide detail")
        return original(value)

    monkeypatch.setattr(pptx_module, "normalized_text", fail_one)
    result = PPTXAdapter().extract(BytesIO(output.getvalue()), context(multiformat_limits))
    assert result.outcome == "PARTIAL"
    assert [unit.status for unit in result.units] == ["FAILED", "USABLE"]
    assert result.units[0].issue_code == "PPTX_SLIDE_ERROR"
    assert "private slide detail" not in repr(result.diagnostics)
    assert_canonical_document_contract(result)


class ConstantEmbeddingProvider:
    provider_name = "deterministic"
    model_name = "deterministic-multiformat"
    dimension = EMBEDDING_DIMENSION

    @staticmethod
    def _vector() -> list[float]:
        return [1.0] + [0.0] * (EMBEDDING_DIMENSION - 1)

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        return [self._vector() for _ in texts]

    def embed_query(self, text: str) -> list[float]:
        return self._vector()


class GroundedTestLLM:
    provider_name = "deterministic"
    model_name = "deterministic-grounded"

    def generate(self, system_instruction: str, user_prompt: str) -> GenerationResult:
        assert "AUTHORIZED_EVIDENCE" in user_prompt
        return GenerationResult(
            "The answer is grounded across the heterogeneous source collection.",
            self.model_name,
            "STOP",
        )


@pytest.fixture
def multiformat_api(tmp_path: Path):
    storage = LocalFileStorage(tmp_path / "multiformat-private-storage")
    embedding = ConstantEmbeddingProvider()
    llm = GroundedTestLLM()
    app.dependency_overrides[get_file_storage] = lambda: storage
    app.dependency_overrides[get_embedding_provider] = lambda: embedding
    app.dependency_overrides[get_llm_provider] = lambda: llm
    app.dependency_overrides[get_ocr_provider] = lambda: object()
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.pop(get_file_storage, None)
        app.dependency_overrides.pop(get_embedding_provider, None)
        app.dependency_overrides.pop(get_llm_provider, None)
        app.dependency_overrides.pop(get_ocr_provider, None)


def _authenticated_workspace(client: TestClient) -> dict:
    registration = client.post(
        "/api/v1/auth/register",
        json={
            "email": "heterogeneous-rag@example.com",
            "password": "heterogeneous-password",
            "display_name": "Heterogeneous researcher",
        },
    )
    assert registration.status_code == 201
    login = client.post(
        "/api/v1/auth/login",
        json={
            "email": "heterogeneous-rag@example.com",
            "password": "heterogeneous-password",
        },
    )
    assert login.status_code == 200
    workspace = client.post(
        "/api/v1/workspaces", json={"name": "Multi-format knowledge base"}
    )
    assert workspace.status_code == 201
    return workspace.json()


def test_all_formats_share_upload_processing_index_retrieval_rag_and_citations(
    multiformat_api: TestClient,
):
    workspace = _authenticated_workspace(multiformat_api)
    sources = [
        ("evidence.pdf", build_text_pdf(["PDF evidence for Aurora"]), "application/pdf"),
        ("evidence.txt", b"TXT evidence for Aurora", "text/plain"),
        ("evidence.md", b"# Markdown evidence\nAurora research", "text/markdown"),
        ("evidence.csv", b"source,finding\nCSV,Aurora evidence\n", "text/csv"),
        (
            "evidence.tsv",
            b"source\tfinding\nTSV\tAurora evidence\n",
            "text/tab-separated-values",
        ),
        (
            "evidence.docx",
            word_bytes(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        (
            "evidence.xlsx",
            workbook_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        (
            "evidence.pptx",
            presentation_bytes(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
    ]
    expected_locator_types = {
        "evidence.pdf": "pdf_page",
        "evidence.txt": "text_range",
        "evidence.md": "markdown_section",
        "evidence.csv": "delimited_rows",
        "evidence.tsv": "delimited_rows",
        "evidence.docx": "docx_paragraph",
        "evidence.xlsx": "xlsx_range",
        "evidence.pptx": "pptx_slide",
    }
    document_ids: dict[str, UUID] = {}

    for filename, data, media_type in sources:
        uploaded = multiformat_api.post(
            f"/api/v1/workspaces/{workspace['id']}/documents",
            files={"file": (filename, data, media_type)},
        )
        assert uploaded.status_code == 201, uploaded.text
        document_id = UUID(uploaded.json()["id"])
        document_ids[filename] = document_id
        processed = multiformat_api.post(
            f"/api/v1/workspaces/{workspace['id']}/documents/{document_id}/process"
        )
        assert processed.status_code == 200, processed.text
        assert processed.json()["status"] == "PROCESSED"
        indexed = multiformat_api.post(
            f"/api/v1/workspaces/{workspace['id']}/documents/{document_id}/index"
        )
        assert indexed.status_code == 200, indexed.text
        assert indexed.json()["indexing_status"] == "INDEXED"

    with SessionLocal() as db:
        for filename, document_id in document_ids.items():
            active_run = db.scalar(
                select(IngestionRun).where(
                    IngestionRun.document_id == document_id,
                    IngestionRun.status == "ACTIVE",
                )
            )
            assert active_run is not None
            first_usable = db.scalar(
                select(ContentUnit)
                .where(
                    ContentUnit.ingestion_run_id == active_run.id,
                    ContentUnit.status == "USABLE",
                )
                .order_by(ContentUnit.ordinal)
            )
            assert first_usable is not None
            assert first_usable.locator["type"] == expected_locator_types[filename]
            stored = db.get(Document, document_id)
            assert stored is not None and stored.indexing_status == "INDEXED"

    search = multiformat_api.post(
        f"/api/v1/workspaces/{workspace['id']}/search",
        json={"query": "Aurora evidence", "top_k": 8},
    )
    assert search.status_code == 200, search.text
    assert {item["filename"] for item in search.json()["items"]} == {
        filename for filename, _, _ in sources
    }

    answer = multiformat_api.post(
        f"/api/v1/workspaces/{workspace['id']}/rag",
        json={"query": "What evidence is present across the collection?", "top_k": 8},
    )
    assert answer.status_code == 200, answer.text
    payload = answer.json()
    assert payload["insufficient_evidence"] is False
    assert {source["filename"] for source in payload["sources"]} == {
        filename for filename, _, _ in sources
    }
    assert {
        source["filename"]: source["locator"]["type"]
        for source in payload["sources"]
    } == expected_locator_types
